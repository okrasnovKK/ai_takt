"""Pulse Survey Service for ai-takt — quarterly engagement pulse survey.

Manages the lifecycle of pulse surveys:
  - created → active (announced + sent to users) → completed

Data is persisted to data/pulse_surveys.json to survive bot restarts.
Answers are stored ANONYMOUSLY — no user_id, username, or chat_id
in the output Excel file.

States per user:
  - survey_sent: questions sent, awaiting answers
  - survey_completed: user answered or skipped
  - reminded: reminder sent (once per survey)
"""
import json
import logging
import re
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

# ── Paths ──
DATA_DIR = Path(__file__).parent.parent / "data"
SURVEYS_STATE_PATH = DATA_DIR / "pulse_surveys.json"
SURVEY_EXCEL_DIR = Path(
    r"C:\Users\User\Ouroboros\Deliverables\excel_data"
)

# ── Survey questions (pulse / engagement) ──
PULSE_QUESTIONS = [
    ("1", "Что в твоей текущей работе приносит тебе наибольшее удовлетворение?"),
    ("2", "Что бы ты хотел улучшить в рабочих процессах или в жизни команды?"),
    ("3", "Любые другие мысли, которыми хочешь поделиться (совсем коротко)?"),
]

PULSE_QUESTION_KEYS = ["q1", "q2", "q3"]


# ══════════════════════════════════════════════
#  Text templates
# ══════════════════════════════════════════════

def get_quarter_label(dt: Optional[datetime] = None) -> str:
    """Return the quarter label, e.g. 'Q3_2026'."""
    if dt is None:
        dt = datetime.now()
    quarter = (dt.month - 1) // 3 + 1
    return f"Q{quarter}_{dt.year}"


def get_quarter_display(dt: Optional[datetime] = None) -> str:
    """Return human-readable quarter, e.g. 'III квартал 2026'."""
    if dt is None:
        dt = datetime.now()
    quarter = (dt.month - 1) // 3 + 1
    roman = ["I", "II", "III", "IV"][quarter - 1]
    return f"{roman} квартал {dt.year}"


def get_pulse_survey_intro() -> str:
    """Return the pulse survey introduction for DM."""
    q_display = get_quarter_display()
    return (
        f"\U0001f331 Привет, команда! Раз в квартал я провожу мини-опрос, "
        f"чтобы узнать ваше настроение и идеи.\n\n"
        f"Этот опрос — за {q_display}. Ответы анонимны, "
        f"я вижу только общую статистику.\n"
        f"Ты можешь не отвечать — это не повлияет на твою работу.\n\n"
        f"Просто напиши ответы в одном сообщении "
        f"(или частями). Если не хочешь отвечать — "
        f"напиши «пропуск»."
    )


def get_pulse_survey_questions_text() -> str:
    """Return the formatted list of pulse survey questions."""
    lines = ["\U0001f4cb <b>Три вопроса:</b>\n"]
    for num, question in PULSE_QUESTIONS:
        lines.append(f"<b>{num}.</b> {question}")
    lines.append(
        "\nНапиши ответы по порядку (например: "
        "«1. Работа в команде 2. Больше созвонов 3. Всë хорошо»)"
        " или частями. Если не хочешь отвечать — "
        "напиши «пропуск»."
    )
    return "\n".join(lines)


def get_pulse_announce_text() -> str:
    """Return the group chat announcement text."""
    q_display = get_quarter_display()
    return (
        f"\U0001f331 <b>Опрос вовлечённости — {q_display}</b>\n\n"
        f"Всем привет! Я запускаю квартальный анонимный опрос, "
        f"чтобы узнать ваше настроение и идеи.\n\n"
        f"Я разошлю короткие вопросы в личные сообщения каждому из вас.\n"
        f"Пожалуйста, найдите минутку и ответьте — "
        f"это поможет сделать нашу команду ещё лучше!\n\n"
        f"Ответы анонимны \u2705\n"
        f"Если не хотите отвечать — просто напишите «пропуск».\n\n"
        f"Если вы ещё не общались с ботом — напишите мне /start, "
        f"чтобы я мог отправить вам опрос."
    )


def get_pulse_reminder_text() -> str:
    """Return the 24h reminder text."""
    return (
        "\u23f0 Напоминаю! Ты ещё не ответил на "
        "квартальный опрос вовлечённости.\n\n"
        "Это займёт всего минуту — просто напиши ответы на "
        "3 вопроса в этом чате. Если не хочешь — напиши «пропуск»."
    )


def get_pulse_completion_text() -> str:
    """Return the thank-you message after survey completion."""
    return (
        "\u2705 Спасибо, твой ответ принят! "
        "Результаты опроса будут переданы организаторам анонимно."
    )


def get_pulse_missing_questions_text(answers: dict) -> str:
    """Return text of unanswered questions to ask again."""
    unanswered = [k for k in PULSE_QUESTION_KEYS if not answers.get(k)]
    if not unanswered:
        return ""

    q_map = dict(zip(PULSE_QUESTION_KEYS, PULSE_QUESTIONS))
    lines = ["\U0001f4cb <b>Остались вопросы:</b>\n"]
    for key in unanswered:
        num, question = q_map[key]
        lines.append(f"<b>{num}.</b> {question}")
    return "\n".join(lines)


def get_pulse_summary_text(survey_data: dict) -> str:
    """Return a summary of survey results for the group chat."""
    q_display = get_quarter_display()
    participants = survey_data.get("participants", {})
    total = len(participants)
    completed = sum(
        1 for p in participants.values()
        if p.get("state") == "survey_completed"
    )
    return (
        f"\U0001f331 <b>Результаты опроса вовлечённости — {q_display}</b>\n\n"
        f"Всего участников: {total}\n"
        f"Ответили: {completed}\n"
        f"Результаты сохранены в файл.\n\n"
        f"Спасибо всем за участие! \U0001f64c"
    )


# ══════════════════════════════════════════════
#  Answer parsing
# ══════════════════════════════════════════════

def parse_pulse_answers(text: str, existing: dict) -> dict:
    """Parse pulse survey answers from user message.

    Supports:
    - Numbered answers: "1. text 2. text 3. text"
    - Multi-line (one answer per line)
    - Single-line answer to first unanswered question
    - 'пропуск' / 'skip' to skip a question

    Merges with existing answers.
    """
    answers = dict(existing)  # copy existing

    # Check for skip-all
    text_stripped = text.strip().lower()
    if text_stripped in ("пропуск", "пропустить", "skip", "skip all",
                         "не хочу участвовать", "нет", "pass"):
        for key in PULSE_QUESTION_KEYS:
            if key not in answers or not answers[key]:
                answers[key] = ""
        return answers

    # Try numbered answers: "1. text 2. text 3. text"
    parts = re.split(
        r'(?:^|[\s\n])([1-3])\s*[.)\-–\s]\s*',
        text.strip(),
        flags=re.MULTILINE,
    )

    if len(parts) >= 3:
        i = 1
        while i + 1 < len(parts):
            num_str = parts[i].strip()
            answer_text = parts[i + 1].strip()
            if num_str in "123":
                q_key = f"q{num_str}"
                if answer_text and answer_text.lower() not in (
                    "пропуск", "skip", "-", ""
                ):
                    answer_text = re.sub(
                        r'\s*[1-3]\s*[.)\-–\s]\s*$', '', answer_text
                    ).strip()
                    answers[q_key] = answer_text
                else:
                    answers[q_key] = ""
            i += 2
        return answers

    # Multi-line: one answer per line
    lines = [l.strip() for l in text.strip().split("\n") if l.strip()]
    if len(lines) >= 2:
        unanswered = [k for k in PULSE_QUESTION_KEYS if not answers.get(k)]
        for i, line in enumerate(lines):
            if i < len(unanswered):
                if line.lower() not in ("пропуск", "skip", "-", ""):
                    answers[unanswered[i]] = line
                else:
                    answers[unanswered[i]] = ""
        return answers

    # Single line — assign to first unanswered question
    unanswered = [k for k in PULSE_QUESTION_KEYS if not answers.get(k)]
    if unanswered and text_stripped not in ("пропуск", "skip", "-", ""):
        answers[unanswered[0]] = text.strip()
    elif unanswered:
        answers[unanswered[0]] = ""

    return answers


def get_pulse_answered_count(answers: dict) -> int:
    """Return how many questions have been answered (non-empty value)."""
    return sum(1 for k in PULSE_QUESTION_KEYS if answers.get(k))


# ══════════════════════════════════════════════
#  Excel writing
# ══════════════════════════════════════════════

def make_pulse_survey_filename() -> str:
    """Create a filename for the pulse survey results Excel file.

    Format: Опрос_вовлеченности_Q{N}_{Год}.xlsx
    """
    q_label = get_quarter_label()
    return f"Опрос_вовлеченности_{q_label}.xlsx"


def write_pulse_survey_excel(
    participants_answers: dict,
) -> Optional[Path]:
    """Write anonymized pulse survey results to a new Excel file.

    Columns: №, Удовлетворение, Улучшения, Мысли, Дата

    Args:
        participants_answers: dict of {user_id: answers_dict}
            where answers_dict has keys q1, q2, q3

    Returns:
        Path to the created file, or None on error.
    """
    SURVEY_EXCEL_DIR.mkdir(parents=True, exist_ok=True)

    filename = make_pulse_survey_filename()
    filepath = SURVEY_EXCEL_DIR / filename

    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        logger.error("openpyxl not installed")
        return None

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Опрос вовлеченности"

        # ── Header row ──
        headers = [
            "№", "Что приносит удовлетворение", "Что улучшить",
            "Доп. мысли", "Дата",
        ]
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_font = Font(bold=True, color="FFFFFF", size=11)

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        # ── Data rows (anonymized) ──
        row_num = 2
        entry_number = 1
        today_str = datetime.now().strftime("%Y-%m-%d")
        for user_id, answers in participants_answers.items():
            q1 = answers.get("q1", "").strip()
            q2 = answers.get("q2", "").strip()
            q3 = answers.get("q3", "").strip()

            # Skip completely empty entries
            if not q1 and not q2 and not q3:
                continue

            ws.cell(row=row_num, column=1, value=entry_number)
            ws.cell(row=row_num, column=2, value=q1)
            ws.cell(row=row_num, column=3, value=q2)
            ws.cell(row=row_num, column=4, value=q3)
            ws.cell(row=row_num, column=5, value=today_str)

            # Wrap text for readability
            for col in range(1, 6):
                ws.cell(row=row_num, column=col).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

            row_num += 1
            entry_number += 1

        # ── Column widths ──
        col_widths = {1: 6, 2: 50, 3: 50, 4: 50, 5: 14}
        for col, width in col_widths.items():
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(col)
            ].width = width

        # Freeze header row
        ws.freeze_panes = "A2"

        wb.save(filepath)
        logger.info(
            "Saved pulse survey to %s (%d responses)",
            filepath, entry_number - 1,
        )
        return filepath

    except Exception as e:
        logger.error("Failed to write pulse survey Excel: %s", e)
        return None


# ══════════════════════════════════════════════
#  State management
# ══════════════════════════════════════════════

class PulseSurveyState:
    """Persistent pulse survey state manager.

    Stores active surveys, per-user progress, and reminder timestamps.
    Survives bot restarts via JSON persistence.
    """

    def __init__(self, state_path: Optional[Path] = None):
        self.state_path = state_path or SURVEYS_STATE_PATH
        self.state_path.parent.mkdir(parents=True, exist_ok=True)
        self._data: dict = {}
        self._load()

    # ── Persistence ──

    def _load(self) -> None:
        if self.state_path.exists():
            try:
                with open(self.state_path, "r", encoding="utf-8") as f:
                    self._data = json.load(f)
            except (json.JSONDecodeError, Exception):
                logger.warning("Failed to load pulse surveys, resetting")
                self._data = {}
        else:
            self._data = {}

    def _save(self) -> None:
        with open(self.state_path, "w", encoding="utf-8") as f:
            json.dump(self._data, f, ensure_ascii=False, indent=2)

    # ── Survey CRUD ──

    def create_survey(
        self,
        survey_id: str,
        created_by: int,
        created_in_chat: int,
    ) -> dict:
        """Create a new pulse survey.

        Returns the survey data dict.
        """
        survey = {
            "survey_id": survey_id,
            "quarter": get_quarter_label(),
            "created_by": created_by,
            "created_in_chat": created_in_chat,
            "created_at": datetime.utcnow().isoformat(),
            "state": "active",  # active | completed
            "participants": {},
            "output_path": None,
            "completed_at": None,
        }
        self._data[survey_id] = survey
        self._save()
        logger.info(
            "Created pulse survey %s (%s)",
            survey_id, get_quarter_display(),
        )
        return survey

    def get_survey(self, survey_id: str) -> Optional[dict]:
        """Get a survey by its id."""
        return self._data.get(survey_id)

    def get_all_active_surveys(self) -> dict:
        """Get all active (non-completed) surveys."""
        return {
            sid: s for sid, s in self._data.items()
            if s.get("state") == "active"
        }

    def complete_survey(self, survey_id: str, output_path: str) -> None:
        """Mark a survey as completed."""
        survey = self._data.get(survey_id)
        if survey:
            survey["state"] = "completed"
            survey["output_path"] = output_path
            survey["completed_at"] = datetime.utcnow().isoformat()
            self._save()

    # ── Participant state ──

    def add_participant(self, survey_id: str, user_id: int) -> None:
        """Add a user as a participant in the survey."""
        survey = self._data.get(survey_id)
        if not survey:
            return
        uid = str(user_id)
        if uid not in survey["participants"]:
            survey["participants"][uid] = {
                "state": "survey_sent",
                "answers": {},
                "sent_at": datetime.utcnow().isoformat(),
                "reminder_sent": False,
                "reminder_sent_at": None,
                "completed_at": None,
            }
            self._save()

    def set_participant_state(
        self, survey_id: str, user_id: int, state: str
    ) -> None:
        """Set the state of a participant."""
        survey = self._data.get(survey_id)
        if not survey:
            return
        uid = str(user_id)
        if uid in survey["participants"]:
            survey["participants"][uid]["state"] = state
            if state == "survey_completed":
                survey["participants"][uid]["completed_at"] = (
                    datetime.utcnow().isoformat()
                )
            self._save()

    def get_participant_state(
        self, survey_id: str, user_id: int
    ) -> Optional[str]:
        """Get the state of a participant, or None if not participating."""
        survey = self._data.get(survey_id)
        if not survey:
            return None
        uid = str(user_id)
        part = survey["participants"].get(uid)
        return part.get("state") if part else None

    def set_participant_answers(
        self, survey_id: str, user_id: int, answers: dict
    ) -> None:
        """Save answers for a participant."""
        survey = self._data.get(survey_id)
        if not survey:
            return
        uid = str(user_id)
        if uid in survey["participants"]:
            survey["participants"][uid]["answers"] = answers
            self._save()

    def get_participant_answers(
        self, survey_id: str, user_id: int
    ) -> dict:
        """Get saved answers for a participant."""
        survey = self._data.get(survey_id)
        if not survey:
            return {}
        uid = str(user_id)
        part = survey["participants"].get(uid, {})
        return part.get("answers", {})

    # ── Reminder logic ──

    def mark_reminder_sent(self, survey_id: str, user_id: int) -> None:
        """Mark that a reminder has been sent to this user."""
        survey = self._data.get(survey_id)
        if not survey:
            return
        uid = str(user_id)
        part = survey["participants"].get(uid)
        if part:
            part["reminder_sent"] = True
            part["reminder_sent_at"] = datetime.utcnow().isoformat()
            self._save()

    def needs_reminder(self, survey_id: str, user_id: int) -> bool:
        """Check if a user needs a reminder (24h after survey sent).

        Returns True if:
        - User is in 'survey_sent' state
        - Reminder hasn't been sent yet
        - At least 24 hours have passed since survey was sent
        """
        survey = self._data.get(survey_id)
        if not survey:
            return False
        uid = str(user_id)
        part = survey["participants"].get(uid)
        if not part:
            return False

        if part["state"] != "survey_sent":
            return False
        if part.get("reminder_sent"):
            return False

        sent_at_str = part.get("sent_at")
        if not sent_at_str:
            return False
        try:
            sent_at = datetime.fromisoformat(sent_at_str)
            now = datetime.utcnow()
            if now - sent_at >= timedelta(hours=24):
                return True
        except (ValueError, TypeError):
            pass
        return False

    # ── Utility ──

    def get_participants_needing_reminder(
        self, survey_id: str
    ) -> list:
        """Get list of user_ids needing a reminder for this survey."""
        survey = self._data.get(survey_id)
        if not survey:
            return []
        return [
            int(uid) for uid, part in survey["participants"].items()
            if self.needs_reminder(survey_id, int(uid))
        ]

    def get_completed_answers(
        self, survey_id: str
    ) -> dict:
        """Get all completed answers for a survey.

        Returns dict of {user_id: answers_dict} for users who have
        completed the survey. Used for Excel generation.
        """
        survey = self._data.get(survey_id)
        if not survey:
            return {}
        return {
            uid: part["answers"]
            for uid, part in survey["participants"].items()
            if part.get("state") == "survey_completed"
        }

    def get_all_answers(self, survey_id: str) -> dict:
        """Get ALL answers (including partial) for a survey.

        Returns dict of {user_id: answers_dict}.
        """
        survey = self._data.get(survey_id)
        if not survey:
            return {}
        return {
            uid: part["answers"]
            for uid, part in survey["participants"].items()
            if part.get("answers")
        }
